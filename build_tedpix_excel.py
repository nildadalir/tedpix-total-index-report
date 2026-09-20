"""Build a 5-year TEDPIX (Tehran Total Index) Excel report with first/last labels."""

from __future__ import annotations

import json
import re
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import jdatetime
import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INDEX_INS_CODE = "32097828799138957"
INDEX_NAME = "TEDPIX — Tehran Stock Exchange Total Index (شاخص کل)"
TSETMC_PAGE = f"https://www.tsetmc.com/IndexInfo/{INDEX_INS_CODE}"
TSETMC_HISTORY_URLS = [
    f"https://cdn.tsetmc.com/api/Index/GetIndexB2History/{INDEX_INS_CODE}",
    f"http://cdn.tsetmc.com/api/Index/GetIndexB2History/{INDEX_INS_CODE}",
]
TGJU_HISTORY_URL = (
    "https://api.tgju.org/v1/market/indicator/summary-table-data/bourse"
    "?draw=1&start=0&length=5000"
)
TGJU_LIVE_URL = "https://call5.tgju.org/ajax.json"
HF_PARQUET_URL = (
    "https://huggingface.co/api/datasets/Farmaanaa/iran_tedpix_stock_index_daily"
    "/parquet/mart/train/0.parquet"
)

ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "output"
OUTPUT_XLSX = OUTPUT_DIR / "TEDPIX_Total_Index_5Y.xlsx"
LOOKBACK_YEARS = 5
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
    ),
    "Accept": "application/json,text/plain,*/*",
}

NAVY = "1F4E79"
GREEN = "548235"
RED = "C00000"
GOLD = "BF8F00"
LIGHT = "D6EAF8"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
TITLE_FONT = Font(color=NAVY, bold=True, name="Calibri", size=16)
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


def parse_number(value) -> float:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("٬", "").strip()
    text = re.sub(r"<[^>]+>", "", text)
    if text in {"", "-", "nan"}:
        return float("nan")
    return float(text)


def jalali_str(d: date) -> str:
    return jdatetime.date.fromgregorian(date=d).strftime("%Y/%m/%d")


def try_tsetmc() -> pd.DataFrame | None:
    for url in TSETMC_HISTORY_URLS:
        try:
            response = requests.get(url, headers=HEADERS, timeout=8)
            response.raise_for_status()
            payload = response.json()
        except (requests.RequestException, ValueError, json.JSONDecodeError):
            continue
        rows = payload.get("indexB2") or payload.get("indexB2History") or []
        if not rows:
            continue
        records = []
        for row in rows:
            raw_date = row.get("dEven") or row.get("date")
            close = row.get("xNivInuClMresIbs") or row.get("close")
            if raw_date is None or close is None:
                continue
            date_text = str(int(raw_date))
            greg = datetime.strptime(date_text, "%Y%m%d").date()
            records.append(
                {
                    "date": greg,
                    "jalali": jalali_str(greg),
                    "open": parse_number(row.get("xNivInuFMresIbs") or row.get("open")),
                    "high": parse_number(row.get("xNivInuPhMresIbs") or row.get("high")),
                    "low": parse_number(row.get("xNivInuPbMresIbs") or row.get("low")),
                    "close": parse_number(close),
                }
            )
        if records:
            frame = pd.DataFrame(records).drop_duplicates("date").sort_values("date")
            frame.attrs["source"] = f"TSETMC {url}"
            return frame
    return None


def fetch_hf_tsetmc_archive() -> pd.DataFrame:
    """Official TSETMC daily series mirrored on Hugging Face (through last vintage)."""
    cache = ROOT / "data" / "tedpix_tsetmc.parquet"
    cache.parent.mkdir(exist_ok=True)
    try:
        response = requests.get(HF_PARQUET_URL, headers=HEADERS, timeout=60)
        response.raise_for_status()
        cache.write_bytes(response.content)
    except requests.RequestException:
        if not cache.exists():
            raise
    raw = pd.read_parquet(cache)
    records = []
    for row in raw.itertuples(index=False):
        greg = datetime.strptime(str(row.date_greg)[:10], "%Y-%m-%d").date()
        records.append(
            {
                "date": greg,
                "jalali": str(row.period).replace("-", "/"),
                "open": float("nan"),
                "high": parse_number(row.high),
                "low": parse_number(row.low),
                "close": parse_number(row.close),
            }
        )
    frame = pd.DataFrame(records).drop_duplicates("date").sort_values("date")
    frame.attrs["source"] = (
        "TSETMC GetIndexB2History archive (Hugging Face Farmaanaa/iran_tedpix_stock_index_daily)"
    )
    return frame


def fetch_tgju_history() -> pd.DataFrame:
    response = requests.get(TGJU_HISTORY_URL, headers=HEADERS, timeout=30)
    response.raise_for_status()
    payload = response.json()
    records = []
    for row in payload["data"]:
        open_, low, high, close, _chg, _pct, gregorian, jalali = row[:8]
        greg = datetime.strptime(gregorian.replace("/", "-"), "%Y-%m-%d").date()
        records.append(
            {
                "date": greg,
                "jalali": jalali.replace("-", "/"),
                "open": parse_number(open_),
                "high": parse_number(high),
                "low": parse_number(low),
                "close": parse_number(close),
            }
        )
    frame = pd.DataFrame(records).drop_duplicates("date").sort_values("date")
    frame.attrs["source"] = "TGJU bourse history (TSE Total Index reprint)"
    return frame


def fetch_today_last_price() -> dict:
    response = requests.get(TGJU_LIVE_URL, headers=HEADERS, timeout=20)
    response.raise_for_status()
    live = response.json()["current"]["bourse"]
    ts = datetime.strptime(live["ts"], "%Y-%m-%d %H:%M:%S")
    today = ts.date()
    return {
        "date": today,
        "jalali": jalali_str(today),
        "open": parse_number(live.get("p")),
        "high": parse_number(live.get("h")),
        "low": parse_number(live.get("l")),
        "close": parse_number(live.get("p")),
        "timestamp": live["ts"],
        "change": parse_number(live.get("d")),
        "change_pct": parse_number(live.get("dp")),
    }


def is_tse_weekday(d: date) -> bool:
    # TSE trades Saturday–Wednesday (Python: Mon=0 … Sun=6).
    return d.weekday() in {0, 1, 2, 5, 6}


def load_daily() -> tuple[pd.DataFrame, dict, str]:
    today_row = fetch_today_last_price()
    parts = []
    notes = []

    tsetmc = try_tsetmc()
    if tsetmc is not None:
        parts.append(tsetmc)
        notes.append("TSETMC official history API")
    else:
        archive = fetch_hf_tsetmc_archive()
        parts.append(archive)
        notes.append(
            "TSETMC (cdn.tsetmc.com) was unreachable, likely because of VPN. "
            "History through 1405/04/31 came from the TSETMC archive on Hugging Face."
        )
        recent = fetch_tgju_history()
        cutoff = archive["date"].max()
        extra = recent[recent["date"] > cutoff]
        extra = extra[extra["date"].map(is_tse_weekday)]
        if not extra.empty:
            parts.append(extra)
            notes.append(
                "Sessions after that vintage were filled from TGJU's TSE Total Index reprint."
            )

    history = pd.concat(parts, ignore_index=True)
    history["date"] = pd.to_datetime(history["date"]).dt.date
    history = history.drop_duplicates("date", keep="first").sort_values("date")

    today_df = pd.DataFrame(
        [{k: today_row[k] for k in ("date", "jalali", "open", "high", "low", "close")}]
    )
    history = history[history["date"] != today_row["date"]]
    history = pd.concat([history, today_df], ignore_index=True)
    history = history.sort_values("date").reset_index(drop=True)
    notes.append("2026 last close is today's last price from the TGJU live ticker.")
    return history, today_row, " ".join(notes)


def year_bounds(frame: pd.DataFrame, years: list[int]) -> pd.DataFrame:
    rows = []
    for year in years:
        part = frame[frame["year"] == year]
        if part.empty:
            continue
        first = part.iloc[0]
        last = part.iloc[-1]
        change = last["close"] - first["close"]
        rows.append(
            {
                "Year": year,
                "First date": first["date"],
                "First (Jalali)": first["jalali"],
                "First close": first["close"],
                "Last date": last["date"],
                "Last (Jalali)": last["jalali"],
                "Last close": last["close"],
                "Change": change,
                "Change %": change / first["close"] if first["close"] else None,
                "Trading days": int(len(part)),
                "Year high": part["close"].max(),
                "Year low": part["close"].min(),
            }
        )
    return pd.DataFrame(rows)


def jalali_year_bounds(frame: pd.DataFrame, years: list[int]) -> pd.DataFrame:
    work = frame.copy()
    work["jyear"] = work["jalali"].str.slice(0, 4).astype(int)
    rows = []
    for year in years:
        part = work[work["jyear"] == year]
        if part.empty:
            continue
        first = part.iloc[0]
        last = part.iloc[-1]
        change = last["close"] - first["close"]
        rows.append(
            {
                "Jalali year": year,
                "First date": first["date"],
                "First (Jalali)": first["jalali"],
                "First close": first["close"],
                "Last date": last["date"],
                "Last (Jalali)": last["jalali"],
                "Last close": last["close"],
                "Change": change,
                "Change %": change / first["close"] if first["close"] else None,
                "Trading days": int(len(part)),
            }
        )
    return pd.DataFrame(rows)


def style_header(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def autosize(ws, min_width=12, max_width=28) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        length = 0
        for cell in column:
            if cell.value is None:
                continue
            length = max(length, min(len(str(cell.value)), max_width))
        ws.column_dimensions[letter].width = max(min_width, length + 3)


def write_table(ws, start_row: int, frame: pd.DataFrame, date_cols, number_cols, pct_cols) -> int:
    headers = list(frame.columns)
    for col, name in enumerate(headers, 1):
        ws.cell(start_row, col, name)
    style_header(ws, start_row, len(headers))
    for r_idx, record in enumerate(frame.itertuples(index=False), start_row + 1):
        for c_idx, value in enumerate(record, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = THIN
            cell.alignment = Alignment(horizontal="center")
            header = headers[c_idx - 1]
            if header in date_cols and isinstance(value, date):
                cell.number_format = "YYYY-MM-DD"
            elif header in pct_cols and value is not None:
                cell.number_format = "0.00%"
            elif header in number_cols and isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"
            if header in {"Change", "Change %"} and isinstance(value, (int, float)):
                cell.font = Font(color=GREEN if value >= 0 else RED, bold=True)
    return start_row + len(frame)


def style_line_series(series, color: str) -> None:
    # Keep DrawingML valid: one fill, positive width, no marker spPr soup.
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.w = 25000
    series.marker = Marker(symbol="none")


def style_marker_series(series, color: str) -> None:
    series.graphicalProperties.line.noFill = True
    marker = Marker(symbol="diamond", size=10)
    marker.graphicalProperties.solidFill = color
    marker.graphicalProperties.line.solidFill = color
    series.marker = marker
    labels = DataLabelList()
    labels.showVal = True
    labels.showCatName = False
    labels.showSerName = False
    labels.numFmt = "#,##0"
    series.dLbls = labels


def configure_axes(chart: LineChart, x_title: str, y_title: str) -> None:
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    chart.y_axis.numFmt = "#,##0"
    chart.legend.position = "b"
    chart.style = 10
    chart.height = 12
    chart.width = 22


def make_excel_safe_drawing(xml: bytes) -> bytes:
    text = xml.decode("utf-8")
    text = text.replace(
        "<cNvGraphicFramePr />",
        '<cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></cNvGraphicFramePr>',
    )
    text = text.replace(
        "<cNvGraphicFramePr/>",
        '<cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></cNvGraphicFramePr>',
    )
    text = text.replace(
        "<xfrm />",
        '<xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xfrm>',
    )
    text = text.replace(
        "<xfrm/>",
        '<xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xfrm>',
    )
    text = text.replace('cNvPr id="1"', 'cNvPr id="2"')
    text = text.replace('cNvPr id="2" name="Chart 2"', 'cNvPr id="3" name="Chart 2"')
    return text.encode("utf-8")


def sanitize_chart_xml(xml: bytes) -> bytes:
    text = xml.decode("utf-8")
    text = text.replace('<a:ln w="0">', "<a:ln>")
    text = text.replace("<a:ln w=\"0\"/>", "<a:ln><a:noFill/></a:ln>")
    # Excel rejects a line that is both filled and not filled.
    text = text.replace(
        "<a:noFill /><a:solidFill><a:srgbClr val=\"548235\" /></a:solidFill>",
        "<a:noFill/>",
    )
    text = text.replace(
        "<a:noFill /><a:solidFill><a:srgbClr val=\"C00000\" /></a:solidFill>",
        "<a:noFill/>",
    )
    text = text.replace("<a:noFill />", "<a:noFill/>")
    text = text.replace('<a:noFill/><a:prstDash val="solid" />', "<a:noFill/>")
    text = text.replace("<a:noFill/><a:prstDash val=\"solid\"/>", "<a:noFill/>")
    return text.encode("utf-8")


def rewrite_xlsx_parts(path: Path, mutators: dict) -> None:
    buffer = BytesIO()
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(buffer, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename in mutators:
                data = mutators[info.filename](data)
            copied = zipfile.ZipInfo(filename=info.filename, date_time=info.date_time)
            copied.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(copied, data)
    path.write_bytes(buffer.getvalue())


def build_workbook(
    daily: pd.DataFrame,
    greg_summary: pd.DataFrame,
    jalali_summary: pd.DataFrame,
    overlay: pd.DataFrame,
    today_row: dict,
    source_note: str,
    years: list[int],
) -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    wb = Workbook()

    charts = wb.active
    charts.title = "Charts"
    summary = wb.create_sheet("Year first-last")
    daily_ws = wb.create_sheet("Daily")
    overlay_ws = wb.create_sheet("Year overlay data")
    notes = wb.create_sheet("Notes")

    # --- Daily sheet (chart source) ---
    daily_ws["A1"] = "Date"
    daily_ws["B1"] = "Jalali"
    daily_ws["C1"] = "Close"
    daily_ws["D1"] = "First of year"
    daily_ws["E1"] = "Last of year"
    daily_ws["F1"] = "Open"
    daily_ws["G1"] = "High"
    daily_ws["H1"] = "Low"
    daily_ws["I1"] = "Gregorian year"
    style_header(daily_ws, 1, 9)

    first_dates = set(greg_summary["First date"])
    last_dates = set(greg_summary["Last date"])
    for idx, row in enumerate(daily.itertuples(index=False), 2):
        daily_ws.cell(idx, 1, row.date).number_format = "YYYY-MM-DD"
        daily_ws.cell(idx, 2, row.jalali)
        daily_ws.cell(idx, 3, row.close).number_format = "#,##0.00"
        if row.date in first_dates:
            daily_ws.cell(idx, 4, row.close).number_format = "#,##0.00"
        if row.date in last_dates:
            daily_ws.cell(idx, 5, row.close).number_format = "#,##0.00"
        daily_ws.cell(idx, 6, row.open).number_format = "#,##0.00"
        daily_ws.cell(idx, 7, row.high).number_format = "#,##0.00"
        daily_ws.cell(idx, 8, row.low).number_format = "#,##0.00"
        daily_ws.cell(idx, 9, row.year)
        for col in range(1, 10):
            daily_ws.cell(idx, col).border = THIN
    last_daily_row = 1 + len(daily)
    daily_ws.auto_filter.ref = f"A1:I{last_daily_row}"
    daily_ws.freeze_panes = "A2"
    daily_ws.sheet_properties.tabColor = NAVY
    autosize(daily_ws)

    cats = Reference(daily_ws, min_col=1, min_row=2, max_row=last_daily_row)
    data_ref = Reference(daily_ws, min_col=3, min_row=1, max_col=5, max_row=last_daily_row)

    trend = LineChart()
    trend.title = f"TEDPIX Total Index {years[0]}-{years[-1]}"
    configure_axes(trend, "Date", "Index close")
    trend.add_data(data_ref, titles_from_data=True)
    trend.set_categories(cats)
    style_line_series(trend.series[0], NAVY)
    style_marker_series(trend.series[1], GREEN)
    style_marker_series(trend.series[2], RED)

    # --- Overlay sheet ---
    overlay_headers = ["Day of year"] + [str(y) for y in years]
    for col, name in enumerate(overlay_headers, 1):
        overlay_ws.cell(1, col, name)
    style_header(overlay_ws, 1, len(overlay_headers))
    for r_idx, record in enumerate(overlay.itertuples(index=False), 2):
        for c_idx, value in enumerate(record, 1):
            cell = overlay_ws.cell(r_idx, c_idx, None if pd.isna(value) else value)
            if c_idx > 1 and cell.value is not None:
                cell.number_format = "#,##0.00"
            cell.border = THIN
    overlay_last = 1 + len(overlay)
    overlay_ws.freeze_panes = "A2"
    autosize(overlay_ws)

    overlay_chart = LineChart()
    overlay_chart.title = "Each year's trend (aligned by day of year)"
    configure_axes(overlay_chart, "Day of year", "Index close")
    overlay_data = Reference(
        overlay_ws, min_col=2, min_row=1, max_col=1 + len(years), max_row=overlay_last
    )
    overlay_cats = Reference(overlay_ws, min_col=1, min_row=2, max_row=overlay_last)
    overlay_chart.add_data(overlay_data, titles_from_data=True)
    overlay_chart.set_categories(overlay_cats)
    palette = [NAVY, "2E86AB", GOLD, GREEN, RED]
    for series, color in zip(overlay_chart.series, palette):
        style_line_series(series, color)

    # --- Charts landing sheet ---
    charts["A1"] = INDEX_NAME
    charts["A1"].font = TITLE_FONT
    charts.merge_cells("A1:L1")
    charts["A2"] = (
        f"5 calendar years ending {today_row['date'].isoformat()} "
        f"({today_row['jalali']}). 2026 last value is today's last price "
        f"({today_row['close']:,.2f} at {today_row['timestamp']})."
    )
    charts["A2"].font = Font(name="Calibri", size=11, italic=True, color="595959")
    charts.merge_cells("A2:L2")
    charts["A3"] = source_note
    charts["A3"].alignment = Alignment(wrap_text=True)
    charts.merge_cells("A3:L3")
    charts.row_dimensions[3].height = 32

    charts.add_chart(trend, "A5")

    snapshot_row = 26
    charts.cell(snapshot_row, 1, "First and last close of each year")
    charts.cell(snapshot_row, 1).font = Font(color=NAVY, bold=True, size=13)
    charts.merge_cells(start_row=snapshot_row, start_column=1, end_row=snapshot_row, end_column=5)
    headers = ["Year", "First date", "First close", "Last date", "Last close"]
    for col, name in enumerate(headers, 1):
        cell = charts.cell(snapshot_row + 1, col, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
    for offset, rec in greg_summary.iterrows():
        idx = snapshot_row + 2 + int(offset)
        values = [
            rec["Year"],
            rec["First date"],
            rec["First close"],
            rec["Last date"],
            rec["Last close"],
        ]
        for col, value in enumerate(values, 1):
            cell = charts.cell(idx, col, value)
            cell.border = THIN
            cell.alignment = Alignment(horizontal="center")
            if col in {3, 5}:
                cell.number_format = "#,##0.00"
            elif col in {2, 4}:
                cell.number_format = "YYYY-MM-DD"

    charts.add_chart(overlay_chart, "A35")
    charts["A59"] = (
        "Green diamonds = first trading-day close of each Gregorian year. "
        "Red diamonds = last trading-day close (today for 2026). "
        "The second chart overlays the same years by calendar day-of-year."
    )
    charts["A59"].alignment = Alignment(wrap_text=True)
    charts.merge_cells("A59:L59")
    charts.column_dimensions["A"].width = 22
    charts.sheet_properties.tabColor = GOLD
    charts.sheet_view.showGridLines = False

    # --- Summary ---
    summary["A1"] = "First and last close of each year"
    summary["A1"].font = TITLE_FONT
    summary.merge_cells("A1:L1")
    summary["A2"] = (
        f"Gregorian years {years[0]}–{years[-1]}. "
        f"For {years[-1]}, last close is today's last price ({today_row['date'].isoformat()})."
    )
    summary.merge_cells("A2:L2")
    end = write_table(
        summary,
        4,
        greg_summary,
        date_cols={"First date", "Last date"},
        number_cols={"First close", "Last close", "Change", "Year high", "Year low"},
        pct_cols={"Change %"},
    )
    jalali_title_row = end + 3
    summary.cell(jalali_title_row, 1, "Same window grouped by Jalali (Iranian) year")
    summary.cell(jalali_title_row, 1).font = Font(color=NAVY, bold=True, size=13)
    write_table(
        summary,
        jalali_title_row + 1,
        jalali_summary,
        date_cols={"First date", "Last date"},
        number_cols={"First close", "Last close", "Change"},
        pct_cols={"Change %"},
    )
    summary.freeze_panes = "A5"
    summary.row_dimensions[4].height = 28
    autosize(summary, min_width=14, max_width=22)
    summary.sheet_properties.tabColor = GREEN

    # --- Notes ---
    notes["A1"] = "Data notes"
    notes["A1"].font = TITLE_FONT
    bullets = [
        f"Requested TSETMC page: {TSETMC_PAGE}",
        f"Index instrument code: {INDEX_INS_CODE} (شاخص کل / TEDPIX)",
        source_note,
        f"Today's last price source: {TGJU_LIVE_URL} field current.bourse.p",
        f"Today's last price: {today_row['close']:,.2f} on {today_row['date']} {today_row['timestamp']}",
        "First number of each year = close of that year's first trading day.",
        "Last number of each year = close of that year's last trading day; 2026 uses today.",
        "Window is the last 5 Gregorian calendar years, ending today.",
        "Re-run: .venv\\Scripts\\python.exe build_tedpix_excel.py",
    ]
    for idx, text in enumerate(bullets, 3):
        notes.cell(idx, 1, f"• {text}")
        notes.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=8)
        notes.cell(idx, 1).alignment = Alignment(wrap_text=True)
        notes.row_dimensions[idx].height = 22
    notes.column_dimensions["A"].width = 28
    notes.sheet_view.showGridLines = False

    wb.save(OUTPUT_XLSX)
    rewrite_xlsx_parts(
        OUTPUT_XLSX,
        {
            "xl/drawings/drawing1.xml": make_excel_safe_drawing,
            "xl/charts/chart1.xml": sanitize_chart_xml,
            "xl/charts/chart2.xml": sanitize_chart_xml,
        },
    )


def main() -> None:
    history, today_row, source_note = load_daily()
    end_year = today_row["date"].year
    years = list(range(end_year - LOOKBACK_YEARS + 1, end_year + 1))
    start_date = date(years[0], 1, 1)
    daily = history[history["date"] >= start_date].copy().reset_index(drop=True)
    daily["year"] = [d.year for d in daily["date"]]
    daily["doy"] = [d.timetuple().tm_yday for d in daily["date"]]

    greg_summary = year_bounds(daily, years)
    j_start = jdatetime.date.fromgregorian(date=start_date).year
    j_end = jdatetime.date.fromgregorian(date=today_row["date"]).year
    jalali_summary = jalali_year_bounds(daily, list(range(j_start, j_end + 1)))

    overlay = pd.DataFrame({"Day of year": range(1, 367)})
    for year in years:
        part = daily[daily["year"] == year]
        overlay[str(year)] = overlay["Day of year"].map(
            dict(zip(part["doy"], part["close"]))
        )

    build_workbook(daily, greg_summary, jalali_summary, overlay, today_row, source_note, years)
    print(f"Wrote {OUTPUT_XLSX}")
    print(f"Rows: {len(daily)}  Years: {years[0]}-{years[-1]}")
    print(f"Today last: {today_row['close']:,.2f} on {today_row['date']}")
    print(greg_summary.to_string(index=False))


if __name__ == "__main__":
    main()
